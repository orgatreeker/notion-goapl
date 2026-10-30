"""Add California seller side (contractors that lost sign/striping/signal bids), signal strength + reasoning, and a verify-emails tab."""
import sys,json,re,csv,datetime as dt,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
S=sys.argv[1]; F='/home/user/notion-goapl/ca-traffic-sign-buyer-signals-2026-10-30.xlsx'
A={m['id']:m for m in json.load(open(S+'/ca/awards.json'))}
def url(i): m=A[i]; return f"https://{m['slug']}.legistar.com/LegislationDetail.aspx?ID={i}&GUID={m['guid']}"
# Jobs (from council award reports / bid abstracts)
J={
 'RIV8260':dict(buyer='City of Riverside',job='HSIP Cycle 12 Citywide Traffic Signal Controller Replacement, Phase 2 (Bid 8260)',date='2026-08-12',winner='Crosstown Electrical & Data, Inc.',low=2135800.00,kind='Signals',src=url(54460),
   bids=[('Yunex, LLC','Riverside, CA',2213070.00),('Bear Electrical Solutions','Alviso, CA',2301500.00),('Econolite Systems, Inc.','Anaheim, CA',2303298.68),('Elecnor Belco Electric, Inc.','Chino, CA',2313488.00),('Select Electric, Inc.','Anaheim, CA',2366553.00),('Polo Engineering, Inc.','Castaic, CA',2822397.00),('Servitek Electric, Inc.','Norwalk, CA',2949540.00)]),
 'ORA2':dict(buyer='City of Orange',job='Traffic Signal Installation + Fiber, Walnut Ave & Orange St (Bid 26-27.09, re-bid)',date='2026-10-15',winner='Calpromax Engineering, Inc.',low=587975.00,kind='Signals',src=url(5962),
   bids=[('Select Electric, Inc.','Anaheim, CA',606556.00),('Alfaro Communications Construction, Inc.','Compton, CA',616038.18),('Ferreira Construction, Inc.','Rancho Cucamonga, CA',638838.00),('DBX, Inc.','Temecula, CA',647999.00),('Elecnor Belco Electric, Inc.','Chino, CA',653360.00)]),
 'ORA1':dict(buyer='City of Orange',job='Traffic Signal Installation + Fiber, Walnut Ave & Orange St (first round - ALL bids rejected)',date='2026-09-03',winner='None - all bids rejected, re-bid',low=None,kind='Signals',src=url(5870),
   bids=[('Alfaro Communications Construction, Inc.','Compton, CA',542339.79),('Select Electric, Inc.','Anaheim, CA',559585.00),('Elecnor Belco Electric, Inc.','Chino, CA',603536.00),('Servitek Electric, Inc.','Norwalk, CA',606456.00),('DBX, Inc.','Temecula, CA',610973.00),('PTM General Engineering Services, Inc.','Riverside, CA',690019.00)]),
 'HUM1':dict(buyer='Humboldt County',job='HSIP Edgelines on Indian Reservation Road Systems (Contract 223001)',date='2026-06-09',winner='GR Sundberg, Inc.',low=792549.10,kind='Striping',src=url(31037),
   bids=[('Darren Taylor Construction, Inc.','Anderson, CA',979230.00),('Central Striping Service, Inc.','Rancho Cordova, CA',1004930.00)]),
 'HUM2':dict(buyer='Humboldt County',job='HSIP Hammond and Mid-Town Trail Crossings (Contract 594187)',date='2026-07-16',winner='Hooven & Co., Inc.',low=459633.00,kind='Crossings/beacons',src=url(31171),
   bids=[('Mercer-Fraser Company','Eureka, CA',481600.00),('GR Sundberg, Inc.','Arcata, CA',501296.70)]),
 'NAPA':dict(buyer='Napa County',job='HSIP Cycle 11 Silverado Trail Intersection Improvements (RDS 22-26)',date='2026-07-09',winner='FBD Vanguard Construction, Inc.',low=676229.00,kind='Road safety (GC)',src=url(14801),
   bids=[('Zara Construction, Inc. (low but non-responsive)','CA',605179.00),('Mission Constructors Inc.','CA',706421.00),('B&M Civil LLC','CA',721061.00),('Argonaut Constructors','CA',732403.00),('Team Ghilotti, Inc.','CA',789141.00),('W.R. Forde Associates, Inc.','CA',816169.00),('Kerex Engineering','CA',954200.00)]),
 'FON':dict(buyer='City of Fontana',job='Etiwanda Ave at Slover Ave/Airport Dr Intersection Improvement (DE-26-220-SB)',date='2026-07-13',winner='CT&T Concrete Paving Inc.',low=11000774.00,kind='Intersection (GC)',src=url(5766),
   bids=[('Van Engineering Inc.','CA',12435174.50),('Riverside Construction Company, Inc.','Riverside, CA',13866114.90),('Toro Enterprises, Inc.','CA',13885158.50),('Vido Samarzich, Inc.','CA',14878780.00),('Hillcrest Contracting, Inc.','CA',15201991.65)]),
}
EM={ # email, source, phone, confidence
 'Select Electric, Inc.':('estimating@selectelectricinc.com','Company website (selectelectricinc.com)','','High'),
 'Servitek Electric, Inc.':('info@servitekelectric.com','Company website (servitekelectric.com)','','High'),
 'Bear Electrical Solutions':('estimating@bear-electrical.com','Berkeley planholder list, Sep 2024 (Robert Asuncion)','(408) 449-5178','Medium'),
 'GR Sundberg, Inc.':('grs@grsinc.biz','Cal Poly Humboldt pre-bid attendance sheet (Amy Harris)','','Medium'),
 'Ferreira Construction, Inc.':('info@ferreiragroup.com','ferreiraconstruction.com - may be a different (NJ) Ferreira, confirm','','Low'),
 'Zara Construction, Inc. (low but non-responsive)':('info@zaraconstruction.com','zaraconstruction.com - confirm it is the CA firm','','Low'),
 'Mission Constructors Inc.':('info@missionconstructors.com','Company website','','Medium'),
 'Kerex Engineering':('patrick@kerexengineering.com','Company website','','Medium'),
 'Van Engineering Inc.':('info@vaneng.com','Company website (vaneng.com)','','Medium'),
 'Toro Enterprises, Inc.':('estimating@toroenterprises.com','Company website','','High'),
 'PTM General Engineering Services, Inc.':('','No public email found','(951) 710-1000','Phone only'),
 'Elecnor Belco Electric, Inc.':('','No public email found','(909) 993-5470','Phone only'),
 'Alfaro Communications Construction, Inc.':('','No public email found - President Hugo Alfaro, 15614 S Atlantic Ave, Compton','','Address only'),
 'DBX, Inc.':('','No public email found (CSLB #240547)','','None yet'),
}
SPEC=re.compile(r'electric|signal|striping|econolite|yunex|communications|DBX|PTM|polo engineering|mercer-fraser|sundberg|darren taylor',re.I)
GIANT=re.compile(r'Yunex|Econolite',re.I)
losses={}
for k,j in J.items():
    for n,loc,amt in j['bids']: losses.setdefault(n,[]).append(k)
rows=[]
for k,j in J.items():
    for n,loc,amt in j['bids']:
        gap=(amt-j['low'])/j['low'] if j['low'] else None
        days=(dt.date(2026,10,30)-dt.date.fromisoformat(j['date'])).days
        spec=bool(SPEC.search(n)) and 'GC' not in j['kind']
        nl=len(losses[n]); why=[]
        if spec: why.append(f"{j['kind'].lower()} specialist")
        else: why.append('general road contractor (subs out striping/signals)')
        if gap is not None: why.append(f"lost by {gap:.1%}")
        else: why.append('all bids rejected - had to re-bid')
        if nl>1: why.append(f"lost {nl} jobs on this list")
        why.append(f"{days} days ago")
        if GIANT.search(n): why.append('large national firm - lower priority')
        if spec and not GIANT.search(n) and ((gap is not None and gap<0.10) or nl>1 or gap is None): s='Strong'
        elif spec or (gap is not None and gap<0.05): s='Medium'
        else: s='Weak'
        if GIANT.search(n): s='Weak' if s=='Medium' else 'Medium'
        e=EM.get(n,('','Not searched yet','',''))
        rows.append([s,n,loc,j['job'],j['buyer'],dt.datetime.fromisoformat(j['date']),amt,j['low'],j['winner'],None,e[0],e[1],e[2],e[3],'; '.join(why).capitalize()+'.','Not contacted',j['src']])
order={'Strong':0,'Medium':1,'Weak':2}
rows.sort(key=lambda r:(order[r[0]],-r[5].timestamp()))
wb=openpyxl.load_workbook(F)
if 'Sellers (lost bids)' in wb.sheetnames: del wb['Sellers (lost bids)']
ws=wb.create_sheet('Sellers (lost bids)',0)
head=['Signal strength','Losing company','Location','Job they lost','Buyer','Bid date','Their bid ($)','Winning bid ($)','Winner','Gap vs winner','Email','Email source','Phone','Email confidence','Reasoning','Outreach status','Bid record']
ws.append(head)
A_=Font(name='Arial',size=10); H=Font(name='Arial',size=10,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78'); Y=PatternFill('solid',fgColor='FFFF00')
SC={'Strong':'C6EFCE','Medium':'FFF2CC','Weak':'EDEDED'}
for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center')
for i,r in enumerate(rows,start=2):
    r[9]=f'=IF(OR(G{i}="",H{i}=""),"",(G{i}-H{i})/H{i})'; ws.append(r)
    for c in ws[i]: c.font=A_; c.alignment=Alignment(vertical='top',wrap_text=c.column in (4,12,15))
    ws.cell(i,1).fill=PatternFill('solid',fgColor=SC[r[0]]); ws.cell(i,16).fill=Y
    ws.cell(i,6).number_format='mm/dd/yyyy'; ws.cell(i,7).number_format='$#,##0'; ws.cell(i,8).number_format='$#,##0'; ws.cell(i,10).number_format='0.0%'
    c=ws.cell(i,17); c.hyperlink=r[16]; c.value='Open record'; c.font=Font(name='Arial',size=10,color='0563C1',underline='single')
for col,w in zip('ABCDEFGHIJKLMNOPQ',[12,32,18,46,18,11,13,13,26,10,32,34,15,13,50,14,12]): ws.column_dimensions[col].width=w
ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
# Buyer strength + reasoning
def addcols(ws,fn):
    h=[c.value for c in ws[1]]
    if 'Signal strength' in h: return
    ws.insert_cols(1,2)
    ws.cell(1,1,'Signal strength'); ws.cell(1,2,'Reasoning')
    for c in (ws.cell(1,1),ws.cell(1,2)): c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True)
    for r in range(2,ws.max_row+1):
        vals={h[k]:ws.cell(r,k+3).value for k in range(len(h))}
        s,why=fn(vals); ws.cell(r,1,s).fill=PatternFill('solid',fgColor=SC[s]); ws.cell(r,2,why)
        for c in (ws.cell(r,1),ws.cell(r,2)): c.font=A_; c.alignment=Alignment(vertical='top',wrap_text=True)
    ws.column_dimensions['A'].width=12; ws.column_dimensions['B'].width=48
def council(v):
    t=v['Tier (1=hottest)']; st=v['Stage']; d=v['Council date']; days=(dt.datetime(2026,10,30)-d).days
    em='has a contact' if (v.get('Best contact (researched)') or v.get('Email (FCC licence)')) else 'no contact yet'
    if t==1: return 'Strong',f"{st}. Voted {days} days ago, so the job is still ahead and not yet on bid sites. {v['$ in record'] or ''} {em}."
    if t==2: return 'Medium',f"Applied for a state/federal safety grant {days} days ago. Money not won yet (awards expected 2027) - good for an early relationship, not an intro this month. {em}."
    if t==4: return 'Weak',f"Already awarded to a contractor - no buyer opportunity now. Use the winner as a sub/supplier lead instead."
    return 'Weak',f"Early discussion or small item ({st}). Watch it; no money approved yet."
def hsip(v):
    cost=v['Project cost ($)'] or 0
    if cost>=500000: return 'Strong',f"State-funded (HSIP Cycle 12, Feb 2025) ${cost:,.0f} {v['Work type'].lower()} job. Money is approved and most Cycle 12 jobs bid 2026-27. Check the agency's bid page first - some may already be out to bid."
    return 'Medium',f"State-funded (HSIP Cycle 12) but smaller (${cost:,.0f}) {v['Work type'].lower()} job. Good for small local contractors; check if already bid."
addcols(wb['Newest signals (council)'],council); addcols(wb['HSIP funded (Cycle 12)'],hsip)
# Verify tab + CSV
V=[]
for r in rows:
    if r[10]: V.append([r[10],r[1],'Seller (contractor)',r[0],r[14],r[11]])
for name in ('Newest signals (council)','HSIP funded (Cycle 12)'):
    t=wb[name]; h=[c.value for c in t[1]]
    for row in t.iter_rows(min_row=2,values_only=True):
        d=dict(zip(h,row)); e=d.get('Email (FCC licence)')
        buyer=d.get('Buyer') or d.get('Buyer (agency)')
        if e: V.append([e,buyer,'Buyer (agency) - '+name.split(' (')[0],d['Signal strength'],d['Reasoning'],'FCC radio licence ('+(d.get('Email check') or '')+')'])
        best=d.get('Best contact (researched)') or ''
        for e2 in re.findall(r'[\w.+-]+@[\w.-]+\.\w+',best): V.append([e2.lower(),buyer,'Buyer (agency) - researched contact',d['Signal strength'],d['Reasoning'],'Agency website / public record'])
seen=set(); V2=[]
for v in V:
    if v[0].lower() in seen: continue
    seen.add(v[0].lower()); V2.append(v)
V2.sort(key=lambda v:(order[v[3]],v[2]))
if 'Verify emails' in wb.sheetnames: del wb['Verify emails']
vs=wb.create_sheet('Verify emails',1)
vs.append(['Email','Name','Side','Signal strength','Reasoning','Email source','Verification result (paste)'])
for c in vs[1]: c.font=H; c.fill=HF
for v in V2: vs.append(v+[''])
for row in vs.iter_rows(min_row=2):
    for c in row: c.font=A_; c.alignment=Alignment(vertical='top',wrap_text=c.column==5)
    row[3].fill=PatternFill('solid',fgColor=SC[row[3].value]); row[6].fill=Y
for col,w in zip('ABCDEFG',[34,30,30,12,60,36,18]): vs.column_dimensions[col].width=w
vs.freeze_panes='B2'
wb.save(F)
with open('/home/user/notion-goapl/ca-all-emails-2026-10-30.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['Email','Name','Side','Signal strength','Reasoning','Email source']); w.writerows(V2)
from collections import Counter
print('sellers rows',len(rows),Counter(r[0] for r in rows),'unique cos',len(set(r[1] for r in rows)),'with email',len(set(r[1] for r in rows if r[10])))
print('verify',len(V2),Counter((v[2].split(' -')[0],v[3]) for v in V2))
