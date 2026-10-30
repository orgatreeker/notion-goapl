"""Rewrite outreach drafts in the sheet in Saad's connector style (short, plain, one question, signed Gopal)."""
import re,sys,csv,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
f=sys.argv[1]; VCSV=sys.argv[2]
V={x['Email'].strip().lower():x['validation result'] for x in csv.DictReader(open(VCSV,encoding='utf-8-sig'))}
wb=openpyxl.load_workbook(f)
A=Font(name='Arial',size=10)
STOP=set('bids bid estimating estimate estimates est info office sales contact admin mail email main bidding support help accounts service team dept'.split())
DEPT=re.compile(r'dept|department|public|works|highway|hwy|engineer|director|commissioner|supervisor|division|clerk|city|village|county|township|road|manager|fcc|contact|police|chief|office|admin',re.I)
def first_from_email(e):
    if not e: return ''
    loc=e.split('@')[0].lower()
    if '.' in loc:
        p=loc.split('.')[0]
        return p.title() if p.isalpha() and len(p)>=3 and p not in STOP else ''
    return loc.title() if loc.isalpha() and 3<=len(loc)<=5 and loc not in STOP else ''
def first_from_contact(c):
    if not c or DEPT.search(c): return ''
    w=c.replace(',',' ').split()
    return w[0].title() if len(w)>=2 and w[0].isalpha() and len(w[0])>=3 else ''
def hi(n): return f'Hi {n},' if n else 'Hi there,'
def kind(wt):
    wt=(wt or '').lower()
    if 'signal' in wt and 'mark' not in wt: return 'signal'
    if 'signal' in wt: return 'signal and striping'
    if 'sign' in wt and 'mark' not in wt: return 'sign'
    return 'striping and sign'
# --- Contractors ---
seen_co={}
ws=wb['Bid losers']; h=[c.value for c in ws[1]]; ix={k:i+1 for i,k in enumerate(h)}
for r in range(2,ws.max_row+1):
    g=lambda k: ws.cell(r,ix[k]).value
    em=g('Email'); n=first_from_email(em); k=kind(g('Work type'))
    lost=g('Lost bid') or ''; num=lost.split(':')[0].split('-')[-1]
    desc=lost.split(':',1)[1].strip() if ':' in lost else lost
    m=re.search(r'District \d',desc); ag=g('Agency') or ''
    where=m.group(0) if m else (ag.split(' (')[0] if 'local' in ag else ag.replace('IDOT ',''))
    date=g('Bid date'); ds=date.strftime('%b %-d') if date else ''
    win=re.sub(r',?\s*(LLC|Inc|Co|Corp|Corporation|Company|Incorporated)\.?$','',(g('Winner') or '').strip().rstrip('.'),flags=re.I).rstrip(',')
    bid,low=g('Their bid ($)'),g('Low/winning bid ($)')
    gap=(bid-low)/low if bid and low else None
    gapline=f" You were about {max(1,round(gap*100))}% off." if gap is not None and 0<=gap<0.10 else ''
    main=g('Job focus')=='Main work'
    job=('signal' if k=='signal' else ('pavement marking' if 'striping' in k else 'sign')) if main else 'road'
    subj=f"{where} {job} job" if where else f"IDOT contract {num}"
    e1=(f"Subject: {subj}\n\n{hi(n)}\n\nSaw the IDOT {where} {job} job (contract {num}) went to {win} on {ds}.{gapline}\n\n"
        f"I connect Illinois {k} contractors with towns that are getting that work ready. I do this independently.\n\n"
        f"Want me to send you the next town that's buying near you?\n\nGopal")
    e1=re.sub(r'  +',' ',e1)
    ws.cell(r,ix['Email 1']).value=e1
    ws.cell(r,ix['Follow-up 1 (day 3)']).value=f"(reply in same thread)\n\n{('Hi '+n) if n else 'Hi there'}, any thoughts on this?\n\nGopal"
    ws.cell(r,ix['Follow-up 2 (day 7)']).value=f"(reply in same thread)\n\n{('Hi '+n) if n else 'Hi there'}, is there someone else on your team who handles new town work?\n\nGopal"
    co=g('Losing company')
    if co in seen_co:
        for c in ('Email 1','Follow-up 1 (day 3)','Follow-up 2 (day 7)'): ws.cell(r,ix[c]).value=f'(Same company as row {seen_co[co]} - send only that one)'
        continue
    seen_co[co]=r
    if 'Saad' not in (g('Notes') or '') and 'H & H' in (g('Losing company') or '').upper().replace('&',' & ').replace('  ',' ') or 'hh-electric' in (em or ''):
        ws.cell(r,ix['Notes']).value=(g('Notes') or '')+' NOTE: H&H and Louie were used in Saad\'s public video - likely being emailed by many people.'
# --- Towns ---
def town_email(t,contact,email,sig):
    n=first_from_contact(contact) or ''
    short=t.replace('County of ','').strip()+(' County' if t.startswith('County of') else '')
    if t.startswith(('City of','Village of','Township of','Town of')): short=t
    work='traffic signal' if sig and 'signal' in sig.lower() else 'striping and sign'
    crews='signal and intersection work' if work=='traffic signal' else 'road striping and sign replacement'
    e1=(f"Subject: {work.capitalize()} contractor for {short}\n\n{'Hi '+n+',' if n else 'Hi,'}\n\n"
        f"I know a couple of contractors in Illinois with crews available for {crews}.\n\n"
        f"Would you be interested in having another qualified contractor bid on your next project?\n\nGopal")
    f1=f"(reply in same thread)\n\n{'Hi '+n if n else 'Hi'}, any thoughts on this?\n\nGopal"
    f2=f"(reply in same thread)\n\n{'Hi '+n if n else 'Hi'}, is there someone in public works I should send this to instead?\n\nGopal"
    return e1,f1,f2
ws=wb['Towns buying']; h=[c.value for c in ws[1]]; ix={k:i+1 for i,k in enumerate(h)}
for r in range(2,ws.max_row+1):
    g=lambda k: ws.cell(r,ix[k]).value
    e1,f1,f2=town_email(g('Town'),g('Contact on licence'),g('Email'),g('IDOT sign/marking/signal job in 2026'))
    ws.cell(r,ix['Email 1']).value=e1; ws.cell(r,ix['Follow-up 1 (day 3)']).value=f1; ws.cell(r,ix['Follow-up 2 (day 7)']).value=f2
# --- Send today: drop Palatine, refill, new drafts ---
st=wb['Send today']; sh=[c.value for c in st[1]]; si={k:i+1 for i,k in enumerate(sh)}
have=[st.cell(r,si['Town']).value for r in range(2,st.max_row+1)]
towns=[[ws.cell(r,c).value for c in range(1,ws.max_column+1)] for r in range(2,ws.max_row+1)]
T={k:i for i,k in enumerate(h)}
bad=re.compile(r'support|911|kencom|mfd|wastew|communications|police|fire|dispatch|\bpd\b|sheriff',re.I)
for r in range(2,st.max_row+1):
    cur=next((x for x in towns if x[T['Town']]==st.cell(r,si['Town']).value),None)
    if st.cell(r,si['Town']).value=='Village of Palatine' or (cur and str(cur[T['Email check']]).startswith('Other')):
        for t in towns:
            if t[T['Town']] in have or V.get(t[T['Email']])!='VALID' or t[T['Department on licence']]=='Police / fire / EMA' or bad.search((t[T['Email']] or '')+' '+(t[T['Contact on licence']] or '')): continue
            if t[T['Department on licence']]!='Public works / highway' or str(t[T['Email check']]).startswith('Other'): continue
            vals={'Town':t[T['Town']],'Email':t[T['Email']],'Email check result':'VALID','Contact on licence':t[T['Contact on licence']],'Phone':t[T['Phone']],'Why this town':'Public works / highway inbox'}
            for k,v in vals.items(): st.cell(r,si[k]).value=v
            have.append(t[T['Town']]); break
for r in range(2,st.max_row+1):
    t=st.cell(r,si['Town']).value
    row=next(x for x in towns if x[T['Town']]==t)
    e1,f1,f2=town_email(t,row[T['Contact on licence']],row[T['Email']],row[T['IDOT sign/marking/signal job in 2026']])
    st.cell(r,si['Email 1']).value=e1; st.cell(r,si['Follow-up 1 (day 3)']).value=f1; st.cell(r,si['Follow-up 2 (day 7)']).value=f2
    for c in range(1,st.max_column+1): st.cell(r,c).font=A
# --- Playbook tab ---
if 'Saad playbook' in wb.sheetnames: del wb['Saad playbook']
pb=wb.create_sheet('Saad playbook',1)
steps=[('Step','When','Message'),
('1. Contractor opener','Day 1','In "Bid losers" Email 1. Send to sign/striping/signal companies first (Job focus = Main work).'),
('1. Town opener','Day 1, same day','In "Send today" / "Towns buying" Email 1. Say "I\'ve got a ... contractor" only after a contractor replies yes; until then "I know a couple of contractors".'),
('2. Follow-up','Day 3-4, same thread','"Hi {First}, any thoughts on this?" / Gopal'),
('3. Check fit','When both sides reply','Look up the town\'s real project (county board / village agenda, capital plan, Legistar) and the contractor\'s IDOT prequalification and past similar work.'),
('4. Contractor pitch + invoice','After fit check',"""Hi {First},

Thanks for getting back to me.

I've been speaking with the {highway / public works} contact at an Illinois {county / village} that is preparing {project}. {Budget + link to public record.} They asked for a contractor's qualifications, safety record and bonding info before bidding.

I checked your IDOT prequalification and the {jobs they bid} this year, which is why I think you're a strong fit.

My fee is ${fee} flat for the direct introduction. Invoice attached.

Once it's paid, I'll put you in touch with the contact on one email thread, with the project details and what they asked for.

If the opportunity is materially different from what I described, I'll replace the introduction.

Gopal
LinknB2B

(Never name the town before payment.)"""),
('5. Town warm-up','Same day as pitch',"""Hi {Name},

Thanks for explaining what you need for {project}.

I'm an independent connector. I help Illinois contractors find town road projects that match the work they already do.

I've found an Illinois {striping / signal} contractor that bid several IDOT jobs this year and has crews available.

There's no fee or obligation for the {Village / County}. The contractor pays me for the introduction.

I'll connect you both by email so they can send their qualifications, safety record and bonding info.

Gopal
LinknB2B"""),
('6. Introduction','After payment',"""Hi {Town contact}, {Contractor first name},

Introducing you both here.

{Town contact} is with {Village/County}, which is preparing {project}. Their team is reviewing contractor qualifications, safety records and bonding before bidding.

{Contractor first name} is with {Company}. They {what they do} and bid {IDOT jobs} this year.

{Contractor first name}, please send your qualification statement, safety record and bonding info to {town inbox}.

I'll let you both take it from here.

Gopal
LinknB2B"""),
('Rules','Always','Under 80 words, plain text, no links in openers, one question, signed "Gopal". Never promise they will win - you sell a real shot. Skip H&H Electric and Village of Palatine (used in Saad\'s public video).')]
for s in steps: pb.append(list(s))
for c in pb[1]: c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78')
for row in pb.iter_rows(min_row=2):
    for c in row: c.font=A; c.alignment=Alignment(wrap_text=True,vertical='top')
pb.column_dimensions['A'].width=26; pb.column_dimensions['B'].width=22; pb.column_dimensions['C'].width=90
wb.save(f)
print('done')
