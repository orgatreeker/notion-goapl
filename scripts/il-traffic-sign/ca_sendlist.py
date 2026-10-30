"""Build the California send list (Strong/Medium, one email per address) with Saad-style first emails."""
import openpyxl,re,json,datetime as dt
F='ca-traffic-sign-buyer-signals-2026-10-30.xlsx'
wb=openpyxl.load_workbook(F,data_only=True)
def rows(name):
    ws=wb[name]; h=[c.value for c in ws[1]]
    return [dict(zip(h,r)) for r in ws.iter_rows(min_row=2,values_only=True)]
STOP=set('bids bid estimating estimate estimates info office sales contact admin mail grs general'.split())
def first(e):
    loc=e.split('@')[0].lower()
    p=loc.split('.')[0] if '.' in loc else (loc if 3<=len(loc)<=6 else '')
    return p.title() if p.isalpha() and len(p)>=3 and p not in STOP else ''
def nm(c,e=''):
    loc=e.split('@')[0].lower()
    if '.' in loc:
        p=loc.split('.')[0]
        if p.isalpha() and len(p)>=3: return p.title()
    if not c or re.search(r'dept|department|public|works|engineer|director|division|city|county|radio|manager|admin|traffic|clerk|road|fleet|special|services|signal|supervisor|tech|operations|maint|wireless|fcc',c,re.I): return ''
    w=c.replace(',',' ').split()
    if len(w) in (2,3) and w[0].isalpha() and len(w[0])>=3 and w[-1].isalpha() and w[-1].lower() in loc: return w[0].title()
    return ''
BADE=re.compile(r'pd\.|pd@|police|sheriff|fire|911|fcc|licens|radio|wireless|\bit@|itsd|tech|dispatch|ems|school|flysfo|ocde|coe\.|fcle|records|invoice|library',re.I)
out=[];seen=set()
# sellers
for r in rows('Sellers (lost bids)'):
    e=(r['Email'] or '').lower()
    if not e or e in seen or r['Signal strength']=='Weak' or r['Email confidence']=='Low': continue
    seen.add(e); f=first(e)
    d=r['Bid date'].strftime('%b %-d')
    JS={'Walnut':'Walnut Ave & Orange St signal','Controller':'citywide signal controller','Edgelines':'HSIP edgeline striping','Hammond':'Hammond and Mid-Town trail crossings','Silverado':'Silverado Trail intersection','Etiwanda':'Etiwanda/Slover intersection'}
    job=next((v for k,v in JS.items() if k in r['Job they lost']),re.sub(r'\s*\(.*?\)','',r['Job they lost']).strip())
    co=re.sub(r',?\s*(LLC|Inc|Co|Company|Corporation)\.?$','',r['Losing company'].strip().rstrip('.'),flags=re.I).rstrip(',')
    if 'rejected' in (r['Winner'] or ''):
        line=f"Saw the {r['Buyer']} threw out all bids on the {job} job ({d}) and had to re-bid it."
    else:
        g=(r['Their bid ($)']-r['Winning bid ($)'])/r['Winning bid ($)']
        win=re.sub(r',?\s*(LLC|Inc|Co|Company|Corporation)\.?$','',r['Winner'].strip().rstrip('.'),flags=re.I).rstrip(',')
        line=f"Saw the {r['Buyer']} {job} job went to {win} on {d}."+(f" You were about {max(1,round(g*100))}% off." if g<0.10 else '')
    kind='signal' if re.search(r'signal',r['Job they lost'],re.I) else 'striping and safety'
    body=f"Hi {f or 'there'},\n\n{line}\n\nI connect California {kind} contractors with cities and counties that are getting that work ready. I do this independently.\n\nWant me to send you the next agency that's buying near you?\n\nGopal"
    out.append(dict(side='seller',strength=r['Signal strength'],to=e,name=r['Losing company'],subject=f"{r['Buyer']} {kind} job",body=body,why=r['Reasoning']))
# buyers
def buyer_add(e,buyer,strength,why,contact,hook,work):
    e=e.lower()
    if e in seen or BADE.search(e): return
    seen.add(e); n=nm(contact,e) or ''
    crews='signal and intersection work' if 'signal' in work.lower() else 'striping, sign and crosswalk work'
    body=f"Hi{(' '+n) if n else ''},\n\nI know a couple of contractors in California with crews available for {crews}.\n\nWould you be interested in having another qualified contractor bid on {hook}?\n\nGopal"
    subj=('Signal' if 'signal' in work.lower() else 'Striping and sign')+f" contractor for {buyer}"
    out.append(dict(side='buyer',strength=strength,to=e,name=buyer,subject=subj,body=body,why=why))
for r in rows('Newest signals (council)'):
    if r['Signal strength'] not in ('Strong','Medium'): continue
    best=r.get('Best contact (researched)') or ''
    ems=re.findall(r'[\w.+-]+@[\w.-]+\.\w+',best)
    if not ems and r['Email (FCC licence)'] and not str(r['Email check']).startswith('Weak'): ems=[r['Email (FCC licence)']]
    HK={'Fullerton':'the Orangethorpe and Harbor signal project once it goes to bid','Riverside':'the signal backplate job you are re-bidding','Hayward':'the road safety project now in design','Mendocino County':'the road safety work now in design','Fontana':'the Etiwanda and Slover intersection work','Fresno':'the HSIP Cycle 13 safety projects you applied for','Fresno County':'the sign and lighting work in your SS4A application','Rialto':'the HSIP Cycle 13 project you applied for','San Bernardino':'the Arrow Route safety improvements','Sacramento':'the HSIP Cycle 13 projects you applied for'}
    hook=HK.get(r['Buyer'],'your upcoming signal and striping work')
    for e in ems[:1]: buyer_add(e,r['Buyer'],r['Signal strength'],r['Reasoning'],r['Contact on licence'],hook,r['Work type'] or '')
for r in rows('HSIP funded (Cycle 12)'):
    if r['Signal strength'] not in ('Strong','Medium'): continue
    e=r['Email (FCC licence)']
    if not e or str(r['Email check']).startswith('Weak'): continue
    w={'signals/beacons':'signal','striping/markings':'striping','signs':'sign','guardrail':'guardrail'}.get((r['Work type'] or '').split(',')[0].lower(),'safety')
    loc=(r['Location'] or '').strip().rstrip('.')
    loc=loc if len(loc)<=60 and not re.match(r'(various|citywide|county ?wide)',loc,re.I) else ''
    if loc:
        if re.match(r'(The|Intersection|At|On|Along|In|One|Two|Three|Four|Five|Six|Seven|Eight|Nine|Ten|Fifty|Sixty)\b',loc): loc=loc[0].lower()+loc[1:]
        if re.match(r'(on|at|along|in) ',loc): loc=' '+loc
        elif re.match(r'(the intersection|intersection|\d|one|two|three|four|five|six|seven|eight|nine|ten|fifty|sixty)',loc): loc=' at '+loc
        else: loc=' on '+loc
    if (r['Contact on licence'] or '') and re.search(r'\b(sgt|sergeant|captain|lt|lieutenant|officer|chief|deputy)\b',r['Contact on licence'],re.I): continue
    buyer_add(e,r['Buyer (agency)'],r['Signal strength'],r['Reasoning'],r['Contact on licence'],f"your HSIP {w} project"+loc,r['Work type'] or '')
order={'Strong':0,'Medium':1}
out.sort(key=lambda x:(x['side']!='seller',order[x['strength']]))
mon=out[:79]; wed=out[79:]
json.dump(mon,open('ca-sendlist-2026-11-02.json','w'),indent=1); json.dump(wed,open('ca-sendlist-2026-11-04.json','w'),indent=1)
from collections import Counter
print(len(out),Counter((x['side'],x['strength']) for x in out))
for x in [o for o in out if o['side']=='seller'][:2]+[o for o in out if o['side']=='buyer'][:2]: print('---',x['to'],'|',x['subject'],'\n',x['body'])
