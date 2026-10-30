"""Build the California buyer-signals sheet (council agenda signals + HSIP funded projects + FCC town inboxes)."""
import json,re,sys,datetime as dt,openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
S=sys.argv[1]  # scratchpad dir
M=json.load(open(S+'/ca/matters.json')); TW=json.load(open(S+'/fcc/ca_towns.json'))
NOISE=re.compile(r'lot line|PLN\d|shoreline|recruitment|legal services|bond counsel|appointments|notice of completion|recruiting consult|Task Force',re.I)
def stage(t):
    if re.search(r'terminat',t,re.I): return (1,'Contract cancelled - job will be re-bid','Now - next 3 months')
    if re.search(r'appl(y|ication)|submi(t|ttal|ssion)|letters? of support',t,re.I): return (2,'Applying for safety grant (HSIP Cycle 13 / SS4A)','Build 2027-28; design starts after award')
    if re.search(r'professional services|design services|engineering design|geotechnical',t,re.I) and not re.search(r'award of contract|construction contract',t,re.I): return (1,'Design engineer hired - construction bid next','Bid in ~6-12 months')
    if re.search(r'award|awarding|execute an agreement with .* for the .* striping|job order',t,re.I): return (4,'Already awarded - winner may need subs/suppliers','Work starting now')
    return (3,'Early planning / council discussion','6-18 months')
BEST={
 'Fresno County':'Steven E. White, Director of Public Works (stwhite@fresnocountyca.gov, 559-600-4537); Bernard Jimenez, Asst Director (bjimenez@fresnocountyca.gov); Design: DesignServices@fresnocountyca.gov',
 'Fresno':'Jill Gormley, TE, Asst Director - Traffic Engineering (email not public; Public Works 559-621-8650)',
 'Fullerton':'Stephen Bise, Public Works Director (licensed traffic engineer); Traffic Engineering 714-738-6845',
 'Fontana':'Traffic Division via Engineering 909-350-7610',
 'Hayward':'Transportation Division (main 510-583-4000)',
 'Riverside':'Nathan Mustafa, Public Works Director (ex-City Traffic Engineer); Public Works 3900 Main St',
 'Mendocino County':'Dept of Transportation, attn Tyler Franz (707-234-2850)',
 'San Bernardino':'Lynn Merrill, Director of Public Works (909-384-5140)',
 'Sacramento':'Traffic Engineering via 311@cityofsacramento.org; Hector Barron, Traffic Engineer (916-808-2669, may be dated)',
 'Rialto':'Public Works, 150 S. Palm Ave (City Engineer oversees traffic)'}
WEAK=re.compile(r'radio|\bit|itsd|tech|utilit|water|mwd|sbmwd|police|pd@|fire|911|dispatch|comm',re.I)
def echeck(e):
    if not e: return ''
    return 'Weak - radio/IT/utility inbox, use researched contact' if WEAK.search(e.split('@')[0]+' '+e.split('@')[1].split('.')[0]) else 'OK - town inbox (verify)'
# FCC town lookup
def base(n): return re.sub(r'^(City|Village|Town|Township|County) of ','',n).lower().strip()
town={}
for t in TW:
    k=('county:' if t['town'].startswith('County') else '')+base(t['town'])
    pri={'Public works / highway':0,'Town office':1,'Police / fire / EMA':2}[t['dept']]
    if k not in town or pri<town[k][0]: town[k]=(pri,t)
def contact(agency):
    a=agency.strip(); k=('county:'+a[:-7].lower()) if a.lower().endswith(' county') else a.lower()
    v=town.get(k)
    if not v or v[0]==2: return ('','','','')
    t=v[1]; ph=t['phone']; ph=f'({ph[:3]}) {ph[3:6]}-{ph[6:10]}' if len(ph)==10 else ph
    return (t['email'],t['contact'].title() if t['contact'].isupper() else t['contact'],ph,t['dept'])
def work(t):
    w=[]
    if re.search(r'signal|beacon|RRFB|HAWK|backplate',t,re.I): w.append('Signals/beacons')
    if re.search(r'strip|marking|thermoplastic|crosswalk|edgeline|road diet|bike lane',t,re.I): w.append('Striping/markings')
    if re.search(r'\bsign|signage|retroreflect',t,re.I): w.append('Signs')
    if re.search(r'guardrail',t,re.I): w.append('Guardrail')
    if re.search(r'speed hump|traffic calming|speed table|bulb',t,re.I): w.append('Traffic calming')
    return ', '.join(w) or 'Road safety'
money=lambda t:(re.search(r'\$[\d,\.]+(?:\s?(?:million|M))?',t) or [''])[0]
A=Font(name='Arial',size=10); HF=PatternFill('solid',fgColor='1F4E78'); H=Font(name='Arial',size=10,bold=True,color='FFFFFF'); Y=PatternFill('solid',fgColor='FFFF00')
TIERC={1:'C6EFCE',2:'DDEBF7',3:'FFF2CC',4:'EDEDED'}
def sheet(ws,head,rows,widths,link=None,yellow=None,tiercol=None):
    ws.append(head)
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center')
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=A; c.alignment=Alignment(vertical='top',wrap_text=widths.get(c.column,14)>=40)
        if link: 
            cell=row[link-1]
            if cell.value: cell.hyperlink=cell.value; cell.value='Open record'; cell.font=Font(name='Arial',size=10,color='0563C1',underline='single')
        if yellow: row[yellow-1].fill=Y
        if tiercol and isinstance(row[tiercol-1].value,int): row[tiercol-1].fill=PatternFill('solid',fgColor=TIERC[row[tiercol-1].value])
    for i in range(1,len(head)+1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=widths.get(i,14)
    ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
wb=Workbook()
# Tab 1 newest agenda signals
rows=[]
for m in M:
    t=m['title'] or ''
    if NOISE.search(t): continue
    tier,st,when=stage(t)
    e,cn,ph,dept=contact(m['buyer'])
    url=f"https://{m['slug']}.legistar.com/LegislationDetail.aspx?ID={m['id']}&GUID={m['guid']}" if m.get('guid') else ''
    rows.append([tier,m['buyer'],st,work(t),money(t),dt.datetime.fromisoformat(m['date'][:10]),t[:400],when,BEST.get(m['buyer'],''),e,echeck(e),cn,ph,'Not contacted',url])
rows.sort(key=lambda r:(r[0],-r[5].timestamp()))
ws=wb.active; ws.title='Newest signals (council)'
sheet(ws,['Tier (1=hottest)','Buyer','Stage','Work type','$ in record','Council date','What the record says','When work is likely','Best contact (researched)','Email (FCC licence)','Email check','Contact on licence','Phone','Outreach status','Official record'],rows,{2:20,3:34,4:20,6:12,7:70,8:24,9:45,10:30,11:24,12:22,15:13},link=15,yellow=14,tiercol=1)
for r in range(2,ws.max_row+1): ws.cell(r,6).number_format='mm/dd/yyyy'
# Tab 2 HSIP C12
c12=openpyxl.load_workbook(S+'/ca/c12.xlsx',read_only=True)['Cycle12ListFinal']
rows=[]
for r in list(c12.iter_rows(values_only=True))[3:]:
    if not r or not r[4]: continue
    desc=r[10] or ''; w=work(desc)
    if not re.search(r'sign|strip|marking|signal|beacon|RRFB|crosswalk|thermoplastic|edgeline|guardrail|retroreflect|HAWK|delineat|marker',desc,re.I): continue
    e,cn,ph,dept=contact(r[4])
    rows.append([r[4],r[7],w,r[11],r[12],r[13],r[9],desc,e,echeck(e),cn,ph,'Not contacted',r[1]])
rows.sort(key=lambda z:-(z[4] or 0))
ws2=wb.create_sheet('HSIP funded (Cycle 12)')
sheet(ws2,['Buyer (agency)','County','Work type','Project cost ($)','State/federal funds ($)','Category','Location','What will be built','Email (FCC licence)','Email check','Contact on licence','Phone','Outreach status','Caltrans project ID'],rows,{1:22,2:18,3:22,7:40,8:60,9:30,10:24,11:22},yellow=13)
for r in range(2,ws2.max_row+1):
    for c in (4,5): ws2.cell(r,c).number_format='$#,##0'
# Read me
ws3=wb.create_sheet('Read me')
for t in ['California buyer signals - how to use',
 'Newest signals = items California city and county councils voted on or listed since Aug 31, 2026 that involve signals, striping, signs or safety. Pulled from each council\'s public agenda system (Legistar) for 33 cities and 8 counties. These are the earliest public signals - before any bid is posted.',
 'Tier 1 = hottest: design engineer just hired (construction bid next) or a contract was cancelled (re-bid coming). Tier 2 = applied for a state/federal safety grant (money likely 2027). Tier 3 = early planning. Tier 4 = already awarded (winner may need subcontractors or suppliers).',
 'HSIP funded (Cycle 12) = 2025 state list of funded road-safety projects with sign/striping/signal/beacon work. The money is approved; most go out to bid in 2026-27. Source: Caltrans Local Assistance, listoffundedprojectsr120250306.xlsx.',
 'Emails come from each agency\'s FCC radio licence (public works first). Not verified yet - run them through your verification tool before sending.',
 'Sources: webapi.legistar.com (council records), dot.ca.gov HSIP approved project lists, data.fcc.gov ULS (Oct 25, 2026 file).']:
    ws3.append([t]); ws3.cell(ws3.max_row,1).font=A; ws3.cell(ws3.max_row,1).alignment=Alignment(wrap_text=True)
ws3['A1'].font=Font(name='Arial',size=12,bold=True); ws3.column_dimensions['A'].width=120
out='/home/user/notion-goapl/ca-traffic-sign-buyer-signals-2026-10-30.xlsx'; wb.save(out)
t1=wb['Newest signals (council)']; print('newest',t1.max_row-1,'with email',sum(1 for r in t1.iter_rows(min_row=2,values_only=True) if r[9]),'| hsip',ws2.max_row-1,'with email',sum(1 for r in ws2.iter_rows(min_row=2,values_only=True) if r[8]))
from collections import Counter
print(Counter(r[9] for r in ws2.iter_rows(min_row=2,values_only=True)))
