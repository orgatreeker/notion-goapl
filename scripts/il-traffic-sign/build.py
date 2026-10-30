import json,re,datetime as dt
from collections import Counter,defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
S='/tmp/claude-0/-home-user-notion-goapl/3894af6e-7006-5c78-8a81-d4969e11e621/scratchpad/'
C=json.load(open(S+'il/contracts.json')); T=json.load(open(S+'il/tabs.json')); TW=json.load(open(S+'fcc/il_towns.json'))
LID={'09/18/2026':'65d994df-81db-4009-9161-3614364e03e7','07/31/2026':'b60c6247-ee50-416d-a1b5-38fd7949536f','06/12/2026':'6ccbbaa6-54f8-4eb4-b2eb-1359f5db32b5','04/24/2026':'0a2d2a61-d2f8-4c3e-a90b-aa7f56babf4c'}
KW=re.compile(r'pavement marking|striping|thermoplastic|raised reflective|signing|sign structure|sign panel|\bsigns?\b|traffic signal|signal moderniz|signals?\b|pavement markers|recessed reflective|epoxy|flashing yellow',re.I)
PRIM=re.compile(r'^\W*(\(.{0,25}\)\s*)?(traffic signal|taffic signal|signal|pavement marking|thermoplastic|striping|signing|sign |sign$|highway sign|raised reflective|recessed|implementation of flashing|pedestrian signing|epoxy|.{0,60}(signal moderniz|signal install|signal improve|signal upgrade|pavement marking|signing|traffic signals? (at|on|along)))',re.I)
def wtype(d):
    t=[]
    if re.search(r'signal',d,re.I): t.append('Signals')
    if re.search(r'marking|striping|thermoplastic|reflective|markers|epoxy',d,re.I): t.append('Markings')
    if re.search(r'\bsign|signing',d,re.I): t.append('Signs')
    return '/'.join(t) or 'Signs/markings'
FREE=re.compile(r'@(gmail|yahoo|aol|hotmail|outlook|comcast|att|sbcglobal|msn|icloud|live|frontier|mchsi|charter)\.',re.I)
def agency(x):
    m=re.search(r'\(([A-Z][A-Z \.\']+)\)\s*$',x['section'])
    if re.match(r'\d\d-\d{5}-\d\d-',x['section']):
        return (m.group(1).title() if m else x['county']+' County')+' (local agency, let by IDOT)', True
    return f"IDOT District {x['district']}", False
leads=[]; jobs=[]
for x in C:
    if not KW.search(x['desc']): continue
    lid=LID[x['letting']]; num=x['item_contract'].split('-')[-1].upper()
    tab=T.get(f'{lid}|{num}')
    bv={b['vendor']:b for b in x['bidders']}
    PAVE=re.compile(r'overlay|resurfac|milling|reconstruct|cape seal|widening|bridge|patching|HMA|shoulder',re.I)
    prim=bool(PRIM.search(x['desc'])) and not PAVE.search(x['desc'][:120])
    ag,local=agency(x)
    j=dict(x=x,prim=prim,agency=ag,local=local,num=num,tab=tab); jobs.append(j)
    if not tab: continue
    bids=[(v,a) for v,a in tab if isinstance(a,float)]
    if len(bids)<2: continue
    site=[(b['vendor'],b['bid']) for b in x['bidders'] if b['bid']]
    if site:
        wv,wa=site[0]; bids=[z for z in bids if z[0]!=wv]+[(wv,wa)]
    else:
        bids.sort(key=lambda z:z[1]); wv,wa=bids[0]
    bids.sort(key=lambda z:z[1])
    win=bv.get(wv,{}).get('company',wv)
    for v,a in bids:
        if v==wv: continue
        b=bv.get(v)
        if not b: continue
        leads.append(dict(j=j,b=b,bid=a,low=wa,winner=win,rank=[z[0] for z in bids].index(v)+1,nb=len(bids),under=a<wa))
cnt=Counter(l['b']['company'] for l in leads if l['j']['prim'])
cnt_all=Counter(l['b']['company'] for l in leads)
GIANT=re.compile(r'Walsh|Kiewit|Lane Construction|OHLA|Granite|Skanska|Flatiron|Archer Western|Judlau',re.I)
def short(d,n=90):
    d=re.sub(r'^\(.*?\)\s*','',d); return d if len(d)<=n else d[:n].rsplit(' ',1)[0]+'…'
def loc(addr):
    m=re.findall(r'([A-Za-z\.\' ]+), ([A-Z]{2}) \d{5}',addr); return f'{m[-1][0].strip()}, {m[-1][1]}' if m else ''
def drafts(l):
    x=l['j']['x']; gap=(l['bid']-l['low'])/l['low']; co=l['b']['company']; wt=wtype(x['desc']).lower().replace('/',' / ')
    where=f"District {x['district']}" if x['county'].lower().startswith('various') else f"{x['county']} County"
    job=f"IDOT contract {l['j']['num']} ({where}, let {x['letting']})"
    co=re.sub(r',?\s*(LLC|L\.L\.C|Inc|Co|Corp|Corporation|Company|Incorporated)\.?$','',co.strip().rstrip('.'),flags=re.I).strip().rstrip(',')
    e=(f"Subject: {job}\n\nHi {co} team,\n\nSaw your ${l['bid']:,.0f} bid on {job} - {short(x['desc'],70).rstrip('.')}. "
       f"{l['winner']} took it at ${l['low']:,.0f}, {gap:.1%} under you.\n\n"
       f"I run LinknB2B. We connect {wt} contractors with the primes that just won Illinois jobs and still need that work subbed out, "
       f"plus towns with their own sign and striping work coming up.\n\nWant me to send the list of primes and towns that fit your area? No cost to look.\n\n{{Your name}}\nLinknB2B")
    f1=f"Subject: Re: {job}\n\nQuick follow-up - I have the list of primes and towns with {wt} work near {loc(l['b']['address']) or 'you'} ready. Should I send it over?\n\n{{Your name}}"
    f2=f"Subject: Re: {job}\n\nLast note from me. If {wt} subcontract work or town jobs aren't a fit right now, just say so and I'll close this out. If they are, reply 'list' and I'll send it.\n\n{{Your name}}"
    return e,f1,f2
def prio(l):
    gap=(l['bid']-l['low'])/l['low']; p=0; why=[]
    if gap<0: gap=1
    if l['j']['prim']: p+=3; why.append('sign/marking/signal job')
    if gap<0.05: p+=2; why.append(f'lost by {gap:.1%}')
    elif gap<0.10: p+=1
    if cnt_all[l['b']['company']]>=2: p+=1; why.append(f"lost {cnt_all[l['b']['company']]} of these jobs")
    if x_recent(l): p+=1; why.append('Sep letting')
    if GIANT.search(l['b']['company']): p-=3
    return p,why
def x_recent(l): return l['j']['x']['letting']=='09/18/2026'
wb=Workbook(); A=Font(name='Arial',size=10); H=Font(name='Arial',size=10,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78'); Y=PatternFill('solid',fgColor='FFFF00'); LK=Font(name='Arial',size=10,color='0563C1',underline='single')
def sheet(ws,head,rows,widths,link_cols=(),money=(),pct=(),date=(),yellow=()):
    ws.append(head)
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center')
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=A; c.alignment=Alignment(vertical='top',wrap_text=c.column in widths and widths[c.column]>=40)
            if c.column in money: c.number_format='$#,##0'
            if c.column in pct: c.number_format='0.0%'
            if c.column in date: c.number_format='mm/dd/yyyy'
            if c.column in yellow: c.fill=Y
            if c.column in link_cols and c.value: c.hyperlink=c.value[1] if isinstance(c.value,tuple) else c.value
    for i in range(1,len(head)+1): ws.column_dimensions[get_column_letter(i)].width=widths.get(i,16)
    ws.freeze_panes='B2'; ws.auto_filter.ref=ws.dimensions
# Bid losers
ls=sorted(leads,key=lambda l:(-prio(l)[0],(l['bid']-l['low'])/l['low']))
ws=wb.active; ws.title='Bid losers'
head=['Losing company','Location','Lost bid','Agency','Bid date','Their bid ($)','Low/winning bid ($)','Winner','Gap vs low bid','Contact name','Title','Phone','Email','Website','Confidence','Notes','Outreach status','Bid source','Work type','Job focus','Email 1','Follow-up 1 (day 3)','Follow-up 2 (day 7)']
rows=[]
for i,l in enumerate(ls,start=2):
    x=l['j']['x']; b=l['b']; p,why=prio(l); e,f1,f2=drafts(l)
    site='' if not b['email'] or FREE.search(b['email']) else b['email'].split('@')[1]
    note=f"Priority {p}"+(' - '+'; '.join(why) if why else '')+f". Placed {l['rank']} of {l['nb']} bids."
    if not b['email']: note+=' No email on IDOT list - call.'
    if l.get('under'): note+=' Bid was lower than the accepted bid but IDOT did not take it (likely rejected/irregular) - check IDOT page.'
    rows.append([b['company'],loc(b['address']),f"{x['item_contract']}: {short(x['desc'])}",l['j']['agency'],dt.datetime.strptime(x['letting'],'%m/%d/%Y'),l['bid'],l['low'],l['winner'],f'=IF(OR(F{i}="",G{i}=""),"",(F{i}-G{i})/G{i})',
        'Estimating dept (ask for owner)','',b['phone'] if b['phone']!='No Record' else '',b['email'],site,'High' if b['email'] else 'Medium',note,'Not contacted',x['url'],wtype(x['desc']),'Main work' if l['j']['prim'] else 'Part of larger job',e,f1,f2])
sheet(ws,head,rows,{1:30,2:18,3:50,4:26,8:28,13:30,16:45,18:14,21:60,22:45,23:45},link_cols=(18,),money=(6,7),pct=(9,),date=(5,),yellow=(17,))
for r in range(2,ws.max_row+1): ws.cell(r,18).value='Open IDOT page'
# Towns
def base(n): return re.sub(r'^(City|Village|Town|Township|County) of ','',n).lower()
tj=defaultdict(list)
for j in jobs:
    x=j['x']
    if not j['local']: continue
    nm=j['agency'].split(' (')[0]
    key='county:'+nm.lower().replace(' county','') if nm.endswith('County') else nm.lower()
    tj[key].append(f"{x['letting']} {x['item_contract']}: {short(x['desc'],70)}")
for j in jobs:  # towns named in state-job descriptions
    for k,n in re.findall(r'\b(City|Village|Town) of ([A-Z][a-z]+(?: [A-Z][a-z]+)?)',j['x']['desc']):
        tj.setdefault(n.lower(),[]); 
def chk(e,town):
    if FREE.search(e): return 'Free mail - likely an official\'s own inbox'
    d=e.split('@')[1]
    if re.search(r'\.(gov|us)$|\.il\.|twp|township|county|cityof|villageof|vil|town',d,re.I) or any(w in d for w in base(town).replace(' ','').split()): return 'Town domain'
    return 'Other domain - may be the town\'s radio vendor, check'
tw=[]
for t in TW:
    b=base(t['town']); key=('county:'+b) if t['town'].startswith('County') else b
    jobsL=[] if t['town'].startswith('Township') else tj.get(key,[])
    tw.append((0 if jobsL else 1,{'Public works / highway':0,'Town office':1,'Police / fire / EMA':2}[t['dept']],t,jobsL))
tw.sort(key=lambda z:(z[0],z[1],z[2]['town']))
ws2=wb.create_sheet('Towns buying')
head2=['Town','Type','Department on licence','Contact on licence','Phone','Email','Email check','Address','City','ZIP','IDOT sign/marking/signal job in 2026','Licence last updated','Outreach status','Source (FCC licence)','Email 1','Follow-up 1 (day 3)','Follow-up 2 (day 7)']
rows=[]
for _,_,t,jl in tw:
    ph=t['phone']; ph=f'({ph[:3]}) {ph[3:6]}-{ph[6:10]}' if len(ph)==10 else ph
    sig=' | '.join(jl[:3])
    hook=f"I saw {t['town']} had road work with signs or markings let through IDOT this year. " if jl else ''
    e=(f"Subject: Sign, striping and signal contractors for {t['town']}\n\nHi,\n\n{hook}LinknB2B keeps a current list of Illinois sign, pavement marking and traffic signal contractors - "
       f"companies that bid IDOT jobs in the last six months, with their prices on record.\n\nIf {t['town']} has striping, sign replacement or signal work coming up, I can send you 3-5 contractors that work near {t['city']} so you get more bids. No cost to the town.\n\nWho handles that for you?\n\n{{Your name}}\nLinknB2B")
    f1=f"Subject: Re: Sign, striping and signal contractors for {t['town']}\n\nFollowing up - happy to send the contractor list for {t['city']}. Should it go to you or to public works?\n\n{{Your name}}"
    f2=f"Subject: Re: Sign, striping and signal contractors for {t['town']}\n\nLast note. If nothing is coming up this year, no problem - reply 'later' and I'll check back before spring striping season.\n\n{{Your name}}"
    rows.append([t['town'],t['town'].split(' of ')[0],t['dept'],t['contact'].title() if t['contact'].isupper() else t['contact'],ph,t['email'],chk(t['email'],t['town']),t['address'].title() if t['address'].isupper() else t['address'],t['city'],t['zip'],sig,dt.datetime.strptime(t['last_action'],'%Y-%m-%d'),'Not contacted',f"https://wireless2.fcc.gov/UlsApp/UlsSearch/license.jsp?licKey={t['usi']}",e,f1,f2])
sheet(ws2,head2,rows,{1:28,4:24,6:32,7:30,8:28,11:55,14:16,15:60,16:45,17:45},link_cols=(14,),date=(12,),yellow=(13,))
for r in range(2,ws2.max_row+1):
    ws2.cell(r,14).value=f"FCC {TW and ''}licence"
# Jobs
ws3=wb.create_sheet('Jobs')
rows=[]
for j in sorted(jobs,key=lambda j:(not j['prim'],j['x']['letting'])):
    x=j['x']; tab=j['tab'] or []; bids=sorted([a for v,a in tab if isinstance(a,float)])
    bv={b['vendor']:b['company'] for b in x['bidders']}
    site=[(b['company'],b['bid']) for b in x['bidders'] if b['bid']]
    rows.append([x['item_contract'],dt.datetime.strptime(x['letting'],'%m/%d/%Y'),j['agency'],x['county'],wtype(x['desc']),'Main work' if j['prim'] else 'Part of larger job',x['desc'],max(len(bids),1 if site else 0),site[0][1] if site else None,site[0][0] if site else '',x['url']])
sheet(ws3,['Item-contract','Letting','Agency','County','Work type','Job focus','Description','Bids','Accepted bid ($)','Accepted bidder','IDOT page'],rows,{3:30,7:70,10:28,11:14},link_cols=(11,),money=(9,),date=(2,))
for r in range(2,ws3.max_row+1): ws3.cell(r,11).value='Open IDOT page'
ws4=wb.create_sheet('Read me')
for t in ['How to use',
 'Bid losers = companies that bid an Illinois DOT job with sign, pavement marking or traffic signal work (Apr 24 - Sep 18, 2026 lettings) and did not have the low bid. Hottest first (Priority in Notes). "Main work" rows are sign/marking/signal jobs; "Part of larger job" rows are road jobs that include that work.',
 'Emails come from IDOT\'s own authorized-bidder list for that letting, so they are what the company gave IDOT this year.',
 'Bid prices come from IDOT\'s "As Accepted Tabulation of Bids" PDFs (read by OCR - check the IDOT page before quoting a number). Winner and winning price = the bid IDOT accepted, as posted on the IDOT contract page.',
 'Towns buying = every Illinois city, village, town, township and county with an active FCC radio licence updated 2021 or later and an email on file. Towns with an IDOT sign/marking/signal job this year are at the top, then public works / highway inboxes.',
 'Email check: "Other domain" emails may belong to the town\'s radio shop, not the town - confirm before sending.',
 'Yellow cells = update as you work (Outreach status). Replace {Your name} in the drafts.',
 'Sources: webapps1.dot.illinois.gov/WCTB (IDOT Transportation Bulletin), data.fcc.gov/download/pub/uls/complete/l_LMpriv.zip (FCC ULS, file dated Oct 25, 2026).']:
    ws4.append([t]); ws4.cell(ws4.max_row,1).font=A; ws4.cell(ws4.max_row,1).alignment=Alignment(wrap_text=True)
ws4['A1'].font=Font(name='Arial',size=12,bold=True); ws4.column_dimensions['A'].width=120
out='/home/user/notion-goapl/il-traffic-sign-market-2026-10-30.xlsx'; wb.save(out)
print('leads',len(ls),'with email',sum(1 for l in ls if l['b']['email']),'main-work leads',sum(1 for l in ls if l['j']['prim']),'unique cos',len(set(l['b']['company'] for l in ls)),'jobs',len(jobs),'prim jobs',sum(j['prim'] for j in jobs),'jobs w/ tab',sum(1 for j in jobs if j['tab']),'towns',len(tw),'towns w job',sum(1 for z in tw if z[0]==0))
for l in ls[:12]: print(prio(l)[0],l['b']['company'],l['b']['email'],l['j']['x']['item_contract'],f"{(l['bid']-l['low'])/l['low']:.1%}",l['winner'])
